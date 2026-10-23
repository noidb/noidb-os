"""
Wing 등록검토 목록 만들기 — data/rocket-pending.json

"Wing 등록검토" = 쿠팡 Wing(판매자센터)에만 있고 로켓(Supplier Hub)에는 없는 상품.

입력(3개, 3번은 없어도 됨):
  1) 상품공급상태관리 다운로드 xlsx
     — 이 파일에는 로켓 상품과 Wing 상품이 함께 들어 있다(2026-10-23 사용자 확인).
       바코드가 R로 시작하는 줄만 로켓 상품이고, S로 시작하는 줄은 전부 Wing 상품이다.
       → 로켓 판단에는 R 바코드 줄만 쓴다. (예전에는 S 줄까지 로켓으로 봐서 Wing 상품 대부분이 빠졌다.)
  2) 노이드비 상품DB 구글시트를 xlsx로 받은 파일
     — 제품DB·판매중지SKU·실적저조SKU·그 외 재등록SKU·비관리SKU 탭(이름이 _로 시작하는 백업 탭 제외)의
       모델명/품번·모델SKU를 로켓에 이미 있는 모델로 본다.
  3) Wing 상품정보 다운로드 xlsx (Coupang_detailinfo, Template 탭) — 카테고리·모델번호를 보탠다.
  기준 Wing 목록: Wing 가격/재고 다운로드 xlsx (price_inventory) — 업체상품ID(=등록상품ID)·옵션ID·업체상품코드.

사용: python3 scripts/build-rocket-pending.py 공급상태.xlsx 상품DB.xlsx price_inventory.xlsx [detailinfo.xlsx]
skuId는 "wing:옵션ID"로 예전과 같다(사이트에서 삭제·직접 입력한 값이 그대로 이어진다).
"""
import json, re, sys, collections
import openpyxl

supply_path, db_path, price_path = sys.argv[1:4]
detail_path = sys.argv[4] if len(sys.argv) > 4 else ""
norm = lambda s: re.sub(r"[\s\W_]+", "", re.sub(r"노이드비", "", s or "")).lower()
text = lambda v: str(v if v is not None else "").strip()

def model_of(raw):
    """모델번호·업체상품코드에는 상품명이 적힌 경우도 있다 → 모델명 형식(영문 1~4자 + 숫자)만 쓰고, 끝의 색상 코드는 뗀다."""
    m = re.match(r"^([a-z]{1,4}\d{3,7})", text(raw).lower())
    return m.group(1) if m else ""

# 1) 로켓 상품 = 공급상태 파일의 R 바코드 줄
supply = list(openpyxl.load_workbook(supply_path).worksheets[0].iter_rows(min_row=2, values_only=True))
rocket_rows = [r for r in supply if text(r[3]).upper().startswith("R")]
rocket_barcodes = {text(r[3]) for r in rocket_rows}
rocket_names = {norm(text(r[2]).split(" / ")[0]) for r in rocket_rows}

# 2) 로켓 모델 = 구글시트 탭의 모델명/품번·모델SKU
rocket_models = set()
for ws in openpyxl.load_workbook(db_path, read_only=True).worksheets:
    if ws.title.startswith("_"):
        continue
    header = [text(h) for h in next(ws.iter_rows(max_row=1, values_only=True), ())]
    cols = [header.index(name) for name in ("모델명/품번", "모델SKU") if name in header]
    for r in ws.iter_rows(min_row=2, values_only=True):
        for i in cols:
            if i < len(r):
                for part in re.split(r"[\s,/\n]+", text(r[i])):
                    if model_of(part):
                        rocket_models.add(model_of(part))

# 3) Wing 상품정보(있으면) — 등록상품ID별 카테고리, 옵션ID별 모델번호
detail_category, detail_model = {}, {}
if detail_path:
    ws = openpyxl.load_workbook(detail_path)["Template"]
    hdr = [c.value for c in ws[4]]
    i_pid, i_cat, i_opt, i_model = (hdr.index(n) for n in ("등록상품ID", "카테고리", "옵션 ID", "모델번호"))
    for r in ws.iter_rows(min_row=5, values_only=True):
        if r[i_pid]:
            detail_category.setdefault(text(r[i_pid]), text(r[i_cat]))
            if r[i_opt]:
                detail_model[text(r[i_opt])] = text(r[i_model])

# 기준 Wing 목록 — 가격/재고 다운로드(3번째 줄이 머리글)
ws = openpyxl.load_workbook(price_path).worksheets[0]
hdr = [text(c.value) for c in ws[3]]
col = {name: hdr.index(name) for name in ["업체상품 ID", "Product ID", "옵션 ID", "바코드", "업체상품코드", "업체 등록 상품명", "등록 옵션명", "판매상태", "승인상태"]}
rows = []
for r in ws.iter_rows(min_row=4, values_only=True):
    if r[col["업체상품 ID"]]:
        rows.append({k: text(r[i]) for k, i in col.items()})

def wing_model(row):
    return model_of(row["업체상품코드"]) or model_of(detail_model.get(row["옵션 ID"], ""))

# 로켓(R 바코드)에 바코드나 상품명이 있는 Wing 상품 → 그 상품의 모델도 로켓에 있는 모델로 본다.
on_rocket_product = collections.defaultdict(bool)
for row in rows:
    if row["바코드"] in rocket_barcodes or norm(row["업체 등록 상품명"]) in rocket_names:
        on_rocket_product[row["업체상품 ID"]] = True
        if wing_model(row):
            rocket_models.add(wing_model(row))

PREFIX_CATEGORY = {"e": "귀걸이", "n": "목걸이", "r": "반지", "b": "팔찌", "a": "발찌", "p": "피어싱"}
def category_of(model, name, wing_category):
    m = re.match(r"^([wm])([a-z])", model)
    if m and m.group(2) in PREFIX_CATEGORY:
        return PREFIX_CATEGORY[m.group(2)]
    for word in ["발찌", "팔찌", "피어싱", "귀걸이", "반지", "목걸이"]:
        if word in f"{name} {wing_category}":
            return word
    return ""

items, skipped = [], collections.Counter()
seen = set()
for row in rows:
    if on_rocket_product[row["업체상품 ID"]]:
        skipped["로켓에 같은 상품"] += 1
        continue
    model = wing_model(row)
    if model and model in rocket_models:
        skipped["로켓에 같은 모델"] += 1
        continue
    key = f"wing:{row['옵션 ID'] or row['업체상품 ID']}"
    if key in seen:
        continue
    seen.add(key)
    name = row["업체 등록 상품명"]
    items.append({
        "skuId": key, "modelSku": "", "modelName": model,
        "category": category_of(model, name, detail_category.get(row["업체상품 ID"], "")),
        "gender": "남성" if model.startswith("m") or "남성" in name else "여성",
        "productName": name, "optionLabel": row["등록 옵션명"],
        "barcode": row["바코드"], "orderableStatus": row["판매상태"] or row["승인상태"],
        "exposedProductId": row["Product ID"], "wingProductId": row["업체상품 ID"],
        "modelSource": "Wing 모델번호" if model else "모델번호 없음",
    })

sources = [p.split("/")[-1] for p in (supply_path, db_path, price_path, detail_path) if p]
json.dump({"generatedFrom": sources, "items": items},
          open("data/rocket-pending.json", "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
print(len(items), "상품(옵션) ·", len({i["wingProductId"] for i in items}), "등록상품 ·",
      len({i["modelName"] for i in items if i["modelName"]}), "모델 ·",
      dict(collections.Counter(i["modelSource"] for i in items)), "· 뺀 것", dict(skipped))
